#!/usr/bin/env python3
"""Every purchase from every receipt archive, in one spreadsheet.

    python export_purchases.py                    one workbook beside the installs
    python export_purchases.py --provider Amazon  just one provider
    python export_purchases.py --csv              a .csv instead of .xlsx
    python export_purchases.py --root "D:\\path\\to\\installs" --out purchases.xlsx

The receipt apps (Amazon, Target, Walmart, Gap) already record every line
item they saw while downloading, one row per item, in an `<Provider> Order
History.csv` beside the PDFs. Nothing here reads a PDF. This walks the
installs folder, finds those files, including the ones in second-account
folders, and writes them out as one long table with a Provider column, plus
an Orders sheet with one row per order and a Summary sheet with a total per
provider per year.

The workbook is rebuilt from the CSVs every time, so it is never the copy
you edit. Amounts are written as numbers, so Excel can sum them.

Needs `openpyxl` for .xlsx. Without it the tool writes a .csv and says so.
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import OrderedDict
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

HISTORY_SUFFIX = " Order History.csv"
_MONEY = re.compile(r"(-)?\s*\$?\s*(-)?([\d,]*\.?\d+)")

# Output columns, in order. Each maps to the CSV column it comes from, or to
# None when it is computed here.
COLUMNS = [
    ("Provider", None),
    ("Account", "Account Holder"),
    ("Date", "Purchase Date"),
    ("Order Number", "Order or Receipt Number"),
    ("Item", "Item Name"),
    ("Quantity", "Quantity"),
    ("Unit Price", "Unit Price"),
    ("Line Total", "Line Item Total"),
    ("Order Total", "Order Total"),
    ("Type", "Purchase Type"),
    ("Fulfillment", "Fulfillment Method"),
    ("Order Status", "Order Status"),
    ("Return Status", "Return Status"),
    ("Category", "Purchase Summary"),
    ("Receipt PDF", "PDF Filename"),
    ("Processing", "Processing Status"),
    ("Notes", "Notes"),
]
NUMERIC = {"Quantity", "Unit Price", "Line Total", "Order Total"}
MONEY = {"Unit Price", "Line Total", "Order Total"}

ORDER_COLUMNS = ["Provider", "Account", "Date", "Order Number", "Items", "Order Total",
                 "Type", "Fulfillment", "Order Status", "Category", "Receipt PDF"]


def money(text) -> Optional[float]:
    """'$1,234.56' -> 1234.56, '-$5.00' -> -5.0, '' -> None."""
    if text is None:
        return None
    m = _MONEY.search(str(text))
    if not m:
        return None
    value = float(m.group(3).replace(",", ""))
    return -value if (m.group(1) or m.group(2)) else value


def quantity(text) -> Optional[float]:
    if text is None or not str(text).strip():
        return None
    try:
        v = float(str(text).strip())
    except ValueError:
        return None
    return int(v) if v.is_integer() else v


def find_histories(root: Path) -> List[Tuple[str, Path]]:
    """(provider, csv path) for every order-history file one level down.
    Second-account folders hold their own copy and are found the same way."""
    found = []
    if not root.is_dir():
        return found
    for folder in sorted(root.iterdir(), key=lambda p: p.name.lower()):
        if not folder.is_dir() or folder.name.startswith("."):
            continue
        for f in sorted(folder.glob("*" + HISTORY_SUFFIX)):
            found.append((f.name[:-len(HISTORY_SUFFIX)], f))
    return found


def load_purchases(provider: str, path: Path) -> List[dict]:
    """One dict per line item, keyed by the output column names."""
    rows = []
    with open(path, encoding="utf-8-sig", newline="") as fh:
        for raw in csv.DictReader(fh):
            row = {}
            for out, src in COLUMNS:
                value = provider if src is None else (raw.get(src) or "").strip()
                if out in MONEY:
                    value = money(value)
                elif out == "Quantity":
                    value = quantity(value)
                row[out] = value
            if not row["Date"] and not row["Item"] and not row["Order Number"]:
                continue
            rows.append(row)
    return rows


def orders_from(purchases: List[dict]) -> List[dict]:
    """One row per order. The order total repeats on every line item in the
    history, so it is taken once, and Items counts the lines."""
    by_key: "OrderedDict[tuple, dict]" = OrderedDict()
    for p in purchases:
        key = (p["Provider"], p["Account"], p["Order Number"] or p["Date"] + "|" + (p["Receipt PDF"] or ""))
        o = by_key.get(key)
        if o is None:
            o = {c: p.get(c) for c in ORDER_COLUMNS if c != "Items"}
            o["Items"] = 0
            by_key[key] = o
        o["Items"] += 1
        if o.get("Order Total") is None and p.get("Order Total") is not None:
            o["Order Total"] = p["Order Total"]
    return list(by_key.values())


def summary_from(orders: List[dict]) -> List[dict]:
    """Spend per provider per year, from the Orders sheet so each order counts once."""
    totals: Dict[Tuple[str, str], List[float]] = {}
    for o in orders:
        year = (o.get("Date") or "")[:4] or "unknown"
        cell = totals.setdefault((o["Provider"], year), [0, 0.0])
        cell[0] += 1
        cell[1] += o.get("Order Total") or 0.0
    rows = [{"Provider": p, "Year": y, "Orders": n, "Total": round(t, 2)}
            for (p, y), (n, t) in totals.items()]
    rows.sort(key=lambda r: (r["Provider"], r["Year"]), reverse=False)
    return rows


def sort_newest_first(rows: List[dict]) -> List[dict]:
    return sorted(rows, key=lambda r: ((r.get("Date") or ""), r.get("Provider") or "",
                                       r.get("Order Number") or ""), reverse=True)


def write_csv(path: Path, purchases: List[dict]) -> None:
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=[c for c, _ in COLUMNS])
        w.writeheader()
        for p in purchases:
            w.writerow({k: ("" if v is None else v) for k, v in p.items()})


def write_xlsx(path: Path, purchases: List[dict], orders: List[dict], summary: List[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    def sheet(ws, columns, rows, widths):
        ws.append(columns)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([r.get(c) for c in columns])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, col in enumerate(columns, start=1):
            letter = get_column_letter(i)
            ws.column_dimensions[letter].width = widths.get(col, 14)
            if col in MONEY or col == "Total":
                for cell in ws[letter][1:]:
                    cell.number_format = "#,##0.00"

    sheet(wb.active, [c for c, _ in COLUMNS], purchases,
          {"Item": 60, "Order Number": 22, "Category": 26, "Receipt PDF": 44, "Notes": 30,
           "Provider": 11, "Date": 11, "Quantity": 9, "Fulfillment": 14})
    wb.active.title = "Purchases"
    sheet(wb.create_sheet("Orders"), ORDER_COLUMNS, orders,
          {"Order Number": 22, "Category": 26, "Receipt PDF": 44, "Provider": 11, "Date": 11})
    sheet(wb.create_sheet("Summary"), ["Provider", "Year", "Orders", "Total"], summary, {})
    wb.save(path)


def export(root: Path, out: Optional[Path] = None, provider: Optional[str] = None,
           as_csv: bool = False) -> dict:
    """Build the workbook. Returns what was written, for the panel and the CLI."""
    histories = find_histories(root)
    if provider:
        histories = [(p, f) for p, f in histories if p.lower() == provider.lower()]
        if histories:
            provider = histories[0][0]      # the archive's own spelling
    purchases: List[dict] = []
    per_provider: Dict[str, int] = {}
    for prov, f in histories:
        rows = load_purchases(prov, f)
        per_provider[prov] = per_provider.get(prov, 0) + len(rows)
        purchases.extend(rows)
    purchases = sort_newest_first(purchases)
    orders = sort_newest_first(orders_from(purchases))
    summary = summary_from(orders)

    have_openpyxl = True
    if not as_csv:
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            have_openpyxl = False
            as_csv = True
    stem = (provider or "All") + " Purchases"
    if out is None:
        out = root / (stem + (".csv" if as_csv else ".xlsx"))
    elif as_csv and out.suffix.lower() == ".xlsx":
        out = out.with_suffix(".csv")
    out.parent.mkdir(parents=True, exist_ok=True)
    try:
        if as_csv:
            write_csv(out, purchases)
        else:
            write_xlsx(out, purchases, orders, summary)
    except PermissionError:
        # Excel holds the file open with an exclusive lock, and this is the
        # one error a person can fix in a second, so say exactly that.
        raise PermissionError(f"{out.name} is open in another program, Excel most likely. "
                              "Close it and build again.") from None
    return {
        "path": str(out),
        "format": "csv" if as_csv else "xlsx",
        "purchases": len(purchases),
        "orders": len(orders),
        "providers": per_provider,
        "sources": [str(f) for _p, f in histories],
        "openpyxl_missing": not have_openpyxl,
        "written_at": datetime.now().isoformat(timespec="seconds"),
    }


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Every purchase from every receipt archive, in one spreadsheet.")
    ap.add_argument("--root", default=str(Path(__file__).resolve().parent),
                    help="folder holding the installs (default: this file's folder)")
    ap.add_argument("--out", help="where to write (default: '<root>/All Purchases.xlsx')")
    ap.add_argument("--provider", help="one provider only, e.g. Amazon")
    ap.add_argument("--csv", action="store_true", help="write a .csv instead of .xlsx")
    args = ap.parse_args(argv)
    root = Path(args.root).expanduser()
    if not root.is_dir():
        print(f"Not a folder: {root}", file=sys.stderr)
        return 2
    try:
        result = export(root, Path(args.out) if args.out else None, args.provider, args.csv)
    except PermissionError as e:
        print(e, file=sys.stderr)
        return 2
    if not result["sources"]:
        print("No '<Provider> Order History.csv' found under", root)
        print("Receipt apps write one after a run. Statement apps have no line items to export.")
        return 1
    for prov, n in sorted(result["providers"].items()):
        print(f"  {prov:12} {n:6} line items")
    print(f"\n{result['purchases']} purchases across {result['orders']} orders")
    if result["openpyxl_missing"]:
        print("openpyxl is not installed, so this is a .csv. `pip install openpyxl` for .xlsx.")
    print("Wrote", result["path"])
    return 0


if __name__ == "__main__":
    sys.exit(main())
